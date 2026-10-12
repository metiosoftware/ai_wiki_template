#!/usr/bin/env python3
"""
Document Reader: Extracts text, tables, and structured data from
documents (.csv, .txt, .docx, .xlsx, .pdf) in raw/ and outputs clean Markdown.
"""

import sys
import os
import csv
import io
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def convert_csv_to_markdown(file_path: Path) -> str:
    """Read CSV and render formatted Markdown table."""
    try:
        content = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = file_path.read_text(encoding="latin-1")

    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    if not rows:
        return "*Empty CSV file.*"

    header = rows[0]
    md_lines = []
    md_lines.append("| " + " | ".join(h.strip() for h in header) + " |")
    md_lines.append("| " + " | ".join("---" for _ in header) + " |")

    for row in rows[1:]:
        # Pad row if shorter than header
        padded = row + [""] * (len(header) - len(row))
        md_lines.append("| " + " | ".join(cell.strip().replace("\n", " ") for cell in padded[:len(header)]) + " |")

    return "\n".join(md_lines)


def convert_docx(file_path: Path) -> str:
    try:
        import docx
    except ImportError:
        return f"> **Notice**: To parse DOCX files with formatting, install `python-docx` (`pip install python-docx`).\n\nFile reference: `{file_path.name}`"

    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def convert_xlsx(file_path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return f"> **Notice**: To parse Excel XLSX sheets, install `openpyxl` (`pip install openpyxl`).\n\nFile reference: `{file_path.name}`"

    wb = openpyxl.load_workbook(file_path, data_only=True)
    out_parts = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        out_parts.append(f"### Sheet: {sheet_name}\n")
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        out_parts.append("| " + " | ".join(header) + " |")
        out_parts.append("| " + " | ".join("---" for _ in header) + " |")
        for r in rows[1:]:
            cells = [str(c).replace("\n", " ") if c is not None else "" for c in r]
            out_parts.append("| " + " | ".join(cells) + " |")
        out_parts.append("\n")
    return "\n".join(out_parts)


def convert_pdf(file_path: Path) -> str:
    try:
        import pypdf
    except ImportError:
        return f"> **Notice**: To parse PDF documents, install `pypdf` (`pip install pypdf`).\n\nFile reference: `{file_path.name}`"

    reader = pypdf.PdfReader(file_path)
    pages = []
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append(f"<!-- Page {idx + 1} -->\n{text.strip()}")
    return "\n\n---\n\n".join(pages)


def convert_file(file_path: Path, output_dir: Path) -> Path:
    suffix = file_path.suffix.lower()
    stem = file_path.stem

    if suffix in (".csv", ".tsv"):
        body = convert_csv_to_markdown(file_path)
    elif suffix == ".docx":
        body = convert_docx(file_path)
    elif suffix in (".xlsx", ".xls"):
        body = convert_xlsx(file_path)
    elif suffix == ".pdf":
        body = convert_pdf(file_path)
    elif suffix in (".txt", ".md", ".json", ".yaml", ".yml"):
        body = file_path.read_text(encoding="utf-8", errors="replace")
    else:
        body = f"Unsupported file type: {suffix}"

    out_file = output_dir / f"{stem}.extracted.md"
    content = f"# Extracted Content: {file_path.name}\n\n" + body + "\n"
    out_file.write_text(content, encoding="utf-8")
    return out_file


def main():
    if len(sys.argv) < 2:
        print("Usage: python convert_doc.py <path_to_source_file> [output_dir]")
        sys.exit(1)

    in_path = Path(sys.argv[1])
    if not in_path.exists():
        print(f"Error: File '{in_path}' does not exist.", file=sys.stderr)
        sys.exit(1)

    out_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else in_path.parent
    try:
        out_path = convert_file(in_path, out_dir)
        print(f"✓ Converted '{in_path.name}' → '{out_path}'")
        print("Tip: Run /wiki-ingest on this extracted file to organize it into wiki/.")
    except Exception as ex:
        print(f"Error converting document: {ex}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
